"""Regressões de Segurança O3-R; transporte e dados pessoais são sintéticos."""
import asyncio
import copy
import hashlib
import io
import json
import re
import subprocess
import zipfile

import httpx
import openpyxl
import pytest
from sqlalchemy import func, select, delete

from backend.oil.caged import CagedCollector, parse_workbook
from backend.oil.ingest import SourceFailure
from backend.oil.legislation import LawCollector
from backend.oil.models import Coleta, Observacao, ProposicaoSnapshot
from .conftest import ROOT, no_sleep
from .test_o3_caged import entry, ingest_caged
from .test_o3_law import CAMARA, SENADO, ingest_law, payload, replay_law


@pytest.mark.parametrize('digest',[None,'','g'*64,'0'*63,123])
def test_manifest_hash_required_before_download(sf,digest):
    metadata,raw=entry('2026-08');metadata.pop('sha256')
    if digest is not None:metadata['sha256']=digest
    calls=[]
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r:(calls.append(r),httpx.Response(200,content=raw))[1])) as client:
            with pytest.raises(SourceFailure):
                await CagedCollector(sf,client,simulada=True,sleep=no_sleep).collect(metadata)
    asyncio.run(run());assert not calls
    with sf() as session:
        rows=session.scalars(select(Coleta)).all()
        assert rows and all(not r.valida for r in rows)
        assert session.scalar(select(func.count()).select_from(Observacao))==0


@pytest.mark.parametrize('problem',['ratio','total','zero_compressed'])
def test_zip_metadata_checked_before_openpyxl(monkeypatch,problem):
    raw=entry('2026-08')[1];original=zipfile.ZipFile.infolist;opened=[]
    def metadata(archive):
        items=original(archive)
        item=copy.copy(items[0])
        item.file_size=201*1024*1024 if problem=='total' else 10001
        item.compress_size=item.file_size if problem=='total' else 0 if problem=='zero_compressed' else 100
        return [item,*items[1:]]
    monkeypatch.setattr(zipfile.ZipFile,'infolist',metadata)
    def forbidden(*a,**k):opened.append(True);raise AssertionError('Parser acionado antes de validar ZIP.')
    monkeypatch.setattr(openpyxl,'load_workbook',forbidden)
    with pytest.raises(ValueError):parse_workbook(raw,'2026-08',allow_recorte=True)
    assert not opened


def test_small_high_ratio_zip_is_invalid_without_observations(sf):
    stream=io.BytesIO(entry('2026-08')[1])
    with zipfile.ZipFile(stream,'a',compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('synthetic-ratio.bin',b'x'*20000)
    raw=stream.getvalue()
    with pytest.raises(SourceFailure):ingest_caged(sf,raw=raw,override={'sha256':hashlib.sha256(raw).hexdigest()})
    with sf() as session:
        assert not session.scalar(select(Coleta).order_by(Coleta.id.desc())).valida
        assert session.scalar(select(func.count()).select_from(Observacao))==0


@pytest.mark.parametrize('count',[1000,1001])
def test_table_row_limit(sf,count):
    wb=openpyxl.load_workbook(io.BytesIO(entry('2026-08')[1]));ws=wb['Tabela 5.1']
    ws.delete_rows(6,ws.max_row)
    for index in range(count):
        year,month=divmod(2020*12+index,12)
        ws.cell(6+index,2,f'{__import__("backend.oil.caged",fromlist=["MONTHS"]).MONTHS[month]}/{year}')
        ws.cell(6+index,6,1)
    ws.cell(6+count,2,None)
    stream=io.BytesIO();wb.save(stream);raw=stream.getvalue()
    year,month=divmod(2020*12+count-1,12);competencia=f'{year}-{month+1:02d}'
    if count==1000:
        assert len(parse_workbook(raw,competencia))==1000
    else:
        with pytest.raises(SourceFailure):
            ingest_caged(sf,raw=raw,override={'sha256':hashlib.sha256(raw).hexdigest(),'competencia':competencia})
        with sf() as session:
            assert not session.scalar(select(Coleta).order_by(Coleta.id.desc())).valida
            assert session.scalar(select(func.count()).select_from(Observacao))==0


def test_defusedxml_is_enabled_and_pinned():
    assert openpyxl.xml.DEFUSEDXML is True
    assert 'defusedxml==0.7.1' in (ROOT/'backend/requirements.txt').read_text(encoding='utf-8')


def test_no_workbook_or_caged_cutout_in_docs():
    names=subprocess.check_output(['git','ls-files','-z','--','docs'],cwd=ROOT).decode().split('\0')
    assert not [n for n in names if n.endswith('.xlsx') or ('caged' in n.casefold() and 'recorte' in n.casefold())]


def test_versioned_evidence_contains_no_absolute_local_paths():
    names=subprocess.check_output(['git','ls-files','-z','--','docs','frontend/evidencias'],cwd=ROOT).decode().split('\0')
    pattern=re.compile(r'(?<![A-Za-z0-9])[A-Za-z]:\\|/Users/|/home/')
    leaked=[]
    for name in filter(None,names):
        try:text=(ROOT/name).read_text(encoding='utf-8')
        except UnicodeDecodeError:continue  # Bytes comprimidos de PNG não são caminhos de texto.
        if pattern.search(text):leaked.append(name)
    assert not leaked


def test_deputy_storage_is_minimized_but_hash_tracks_original(sf,client):
    synthetic={'dados':{'id':1,'nomeCivil':'Nome sintético','cpf':'CPF_SINTETICO_INVALIDO',
        'dataNascimento':'DATA_SINTETICA_INVALIDA','email':'teste@example.invalid',
        'endereco':{'rua':'Endereco sintetico'},'ultimoStatus':{'nome':'Nome sintético','siglaPartido':'TESTE','siglaUf':'DF'}}}
    original=httpx.Response(200,json=synthetic).content
    def handler(request):
        if str(request.url)==CAMARA+'/autores':return httpx.Response(200,json={'dados':[{'nome':'Nome sintético','ordemAssinatura':1,'uri':'https://dadosabertos.camara.leg.br/api/v2/deputados/1'}]})
        if '/deputados/1' in str(request.url):return httpx.Response(200,content=original)
        return replay_law()(request)
    ingest_law(sf,handler)
    with sf() as session:
        collection=session.scalar(select(Coleta).where(Coleta.url_consultada.endswith('/deputados/1')))
        saved=json.loads(collection.resposta_bruta)
        assert set(saved)<= {'id','nome','siglaPartido','siglaUf'}
        assert saved['siglaPartido']=='TESTE' and saved['siglaUf']=='DF'
        assert collection.hash_sha256==hashlib.sha256(original).hexdigest()
    corpus=collection.resposta_bruta.decode()+json.dumps(client.get('/api/v1/proposicoes/2233802').json())
    assert all(key not in corpus for key in ('cpf','dataNascimento','email','endereco','CPF_SINTETICO_INVALIDO'))


@pytest.mark.parametrize('declared',[True,False])
def test_law_response_limit_is_502_without_retry_or_snapshot(sf,declared):
    calls=[];source=json.dumps(payload('camara')).encode()
    class Stream(httpx.AsyncByteStream):
        async def __aiter__(self):
            yield source
            for _ in range(6):yield b' '*(1024*1024)
    def handler(request):
        calls.append(request)
        return httpx.Response(200,content=source,headers={'Content-Length':str(5*1024*1024+1)}) if declared else httpx.Response(200,stream=Stream())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            with pytest.raises(SourceFailure) as failure:
                await LawCollector(sf,client,simulada=True,sleep=no_sleep).fetch(2233802,'camara',CAMARA)
            assert (failure.value.status,failure.value.code)==(502,'resposta_invalida')
    asyncio.run(run());assert len(calls)==1
    with sf() as session:
        assert not session.scalar(select(Coleta).order_by(Coleta.id.desc())).valida
        assert session.scalar(select(func.count()).select_from(ProposicaoSnapshot))==0


@pytest.mark.parametrize('url,accepted',[
    ('javascript:alert(1)',False),('http://www.planalto.gov.br/',False),
    ('https://planalto.gov.br.exemplo.com/',False),('https://www.planalto.gov.br/',True),
    ('https://www.in.gov.br/web/dou/',True)])
def test_generated_norm_requires_official_https_source(sf,client,url,accepted):
    senate=copy.deepcopy(payload('senado'));senate['normaGerada']={'rotulo':'Norma sintética de teste','data':'2026-10-16','link':url}
    ingest_law(sf,replay_law(senado=senate));data=client.get('/api/v1/proposicoes/2233802').json()
    assert next(s for s in data['situacao'] if s['casa']=='senado')['disponivel'] is True
    if accepted:assert data['norma_gerada']['link']==url
    else:
        assert data['norma_gerada'] is None
        assert 'fora das fontes oficiais' in data['norma_gerada_motivo']


@pytest.mark.parametrize('missing',[['camara'],['camara','senado']])
def test_valid_collection_without_snapshot_never_returns_500(sf,client,missing):
    ingest_law(sf)
    with sf.begin() as session:
        session.execute(delete(ProposicaoSnapshot).where(ProposicaoSnapshot.casa.in_(missing)))
    response=client.get('/api/v1/proposicoes/2233802')
    assert response.status_code==(200 if len(missing)==1 else 503)
    if len(missing)==1:assert response.json()['situacao_atual']['casa']=='senado'
    else:assert 'ultima_coleta_id' in response.json()['erro']
