"""O3 B4: saldos reais gravados, revisões e reconstrução por competência."""
import asyncio
import hashlib
import importlib
import io
import json
from datetime import datetime, timezone

import httpx
import pytest
from alembic import command
from sqlalchemy import func, select

from backend.oil.ingest import SourceFailure
from backend.oil.models import Coleta, Observacao, Serie
from backend.oil.read import read_series
from .conftest import FIXTURES, alembic_config, no_sleep

ROOT=FIXTURES/'o3'


def entry(month):
    raw=(ROOT/f'caged_t51_{month}_recorte.xlsx').read_bytes()
    return dict(competencia=month,arquivo=f'Fixture {month}.xlsx',drive_file_id='fixture-teste',
        url_download=f'https://drive.usercontent.google.com/download?id=fixture-{month}',
        pagina_oficial='https://www.gov.br/trabalho-e-emprego/pt-br/',publicado_em='2026-09-29',
        sha256=hashlib.sha256(raw).hexdigest(),fixture_recorte=True),raw


def ingest_caged(sf,month='2026-08',*,raw=None,override=None,handler=None):
    metadata,original=entry(month);metadata.update(override or {})
    async def run():
        module=importlib.import_module('backend.oil.caged')
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler or (lambda r:httpx.Response(200,content=original if raw is None else raw)))) as client:
            return await module.CagedCollector(sf,client,simulada=True,sleep=no_sleep).collect(metadata)
    return asyncio.run(run())


def test_o3_fixture_hashes_match_manifest():
    manifest=json.loads((ROOT/'manifesto-fixtures.json').read_text(encoding='utf-8'))
    for item in manifest['fixtures']:
        assert hashlib.sha256((ROOT/item['arquivo'].split('/')[-1]).read_bytes()).hexdigest()==item['sha256']


def test_o3_caged_real_values_attribution_and_annual_difference(sf,client):
    ingest_caged(sf);response=client.get('/api/v1/indicadores/caged-saldo');assert response.status_code==200
    data=response.json();points={p['periodo']:p for p in data['pontos']}
    assert points['202608']['valor']==165827 and points['202512']['valor']==-635078
    assert points['202606']['valor']==136959
    assert points['202608']['variacao_anual']==11144 and points['202607']['variacao_anual']==-77869
    assert data['tipo_valor'] is data['classificacao'] is data['classificacao_origem'] is data['cv'] is None
    assert data['medida']=='saldo' and data['tipo_periodo']=='mensal' and data['unidade']=='postos de trabalho'
    assert data['variacao_anual_origem']=='calculo_observatorio'
    assert data['ajuste']=='com_ajuste' and data['competencia_coleta']=='2026-08'
    assert 'Tabela 5.1 (com ajustes)' in data['atribuicao']['texto'] and 'agosto/2026' in data['atribuicao']['texto']
    assert data['atribuicao']['licenca']=='CC BY-ND 3.0' and data['nota_revisao']
    with sf() as session:
        collection=session.get(Coleta,data['coleta_id']);assert collection.resposta_bruta is None and collection.hash_sha256


@pytest.mark.parametrize('order',[('2026-07','2026-08'),('2026-08','2026-07')])
def test_o3_caged_reconstruct_each_file_even_out_of_order(sf,client,order):
    for month in order:ingest_caged(sf,month)
    reconstruct=importlib.import_module('backend.oil.caged').reconstruct
    with sf() as session:
        july=reconstruct(session,'2026-07',test_mode=True);aug=reconstruct(session,'2026-08',test_mode=True)
        assert july['202607']['valor']==58568 and july['202512']['valor']==-634374
        assert aug['202607']['valor']==55915 and aug['202512']['valor']==-635078 and aug['202608']['valor']==165827
        parse=importlib.import_module('backend.oil.caged').parse_workbook
        for month,points in [('2026-07',july),('2026-08',aug)]:
            assert {p:v['valor'] for p,v in points.items()}==parse(entry(month)[1],month,allow_recorte=True)
    data=client.get('/api/v1/indicadores/caged-saldo').json();assert data['competencia_coleta']=='2026-08'
    assert next(p for p in data['pontos'] if p['periodo']=='202607')['valor']==55915


def test_o3_caged_repeated_file_does_not_create_version(sf):
    ingest_caged(sf,'2026-07');ingest_caged(sf)
    with sf() as session:before=session.scalar(select(func.count()).select_from(Observacao))
    ingest_caged(sf)
    with sf() as session:assert session.scalar(select(func.count()).select_from(Observacao))==before


def test_o3_caged_chronological_revision_versions(sf):
    ingest_caged(sf,'2026-07');ingest_caged(sf)
    with sf() as session:
        sid=session.scalar(select(Serie.id).where(Serie.codigo=='caged-saldo'))
        for period,values in [('202607',[58568,55915]),('202512',[-634374,-635078]),('202605',[81996,81648])]:
            rows=session.scalars(select(Observacao).where(Observacao.serie_id==sid,Observacao.periodo==period).order_by(Observacao.versao)).all()
            assert [(o.versao,int(o.valor)) for o in rows]==list(enumerate(values,1))


@pytest.mark.parametrize('code',['desocupacao','rendimento-real'])
def test_o3_negative_value_trigger_still_blocks_ibge(populated,code):
    from sqlalchemy.exc import IntegrityError
    with populated() as session:
        sid=session.scalar(select(Serie.id).where(Serie.codigo==code))
        collection=session.scalar(select(Coleta).where(Coleta.serie_codigo==code,Coleta.valida.is_(True)))
        session.add(Observacao(serie_id=sid,periodo='201201',valor=-1,cv=None,variacao_anual=None,diferenca_anual=None,classificacao=None,classificacao_origem='calculada',coleta_id=collection.id,versao=1))
        with pytest.raises(IntegrityError):session.commit()


def test_o3_caged_downgrade_preserves_existing_data(sf):
    ingest_caged(sf)
    with pytest.raises(RuntimeError):command.downgrade(alembic_config(),'0003')


def test_o3_caged_returns_to_old_value_creates_new_version(sf):
    import openpyxl
    ingest_caged(sf,'2026-07');ingest_caged(sf)
    workbook=openpyxl.load_workbook(io.BytesIO(entry('2026-08')[1]));sheet=workbook['Tabela 5.1']
    for row in range(6,sheet.max_row+1):
        if sheet.cell(row,2).value=='Julho/2026':sheet.cell(row,6,58568)
    sheet.cell(21,2,'Setembro/2026');sheet.cell(21,6,1)  # Variante sintética, só neste teste.
    stream=io.BytesIO();workbook.save(stream);raw=stream.getvalue()
    metadata,_=entry('2026-08');metadata.update(competencia='2026-09',sha256=hashlib.sha256(raw).hexdigest())
    async def run():
        module=importlib.import_module('backend.oil.caged')
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r:httpx.Response(200,content=raw))) as client:
            return await module.CagedCollector(sf,client,simulada=True,sleep=no_sleep).collect(metadata)
    asyncio.run(run())
    with sf() as session:
        sid=session.scalar(select(Serie.id).where(Serie.codigo=='caged-saldo'))
        assert [int(o.valor) for o in session.scalars(select(Observacao).where(Observacao.serie_id==sid,Observacao.periodo=='202607').order_by(Observacao.versao))]==[58568,55915,58568]


@pytest.mark.parametrize('mutation',['sheet','header','month','hash','integer'])
def test_o3_caged_bad_layout_not_persisted_and_last_good_served(sf,client,mutation):
    import openpyxl
    ingest_caged(sf);metadata,raw=entry('2026-08')
    workbook=openpyxl.load_workbook(io.BytesIO(raw));sheet=workbook['Tabela 5.1']
    if mutation=='sheet':sheet.title='Tabela 5'
    if mutation=='header':sheet.cell(5,6,'Saldo diferente')
    if mutation=='month':sheet.cell(6,2,'mês inválido')
    if mutation=='integer':sheet.cell(6,6,1.5)
    stream=io.BytesIO();workbook.save(stream);raw=stream.getvalue()
    expected='0'*64 if mutation=='hash' else hashlib.sha256(raw).hexdigest()
    with pytest.raises(SourceFailure):ingest_caged(sf,raw=raw,override={'sha256':expected})
    data=client.get('/api/v1/indicadores/caged-saldo').json();assert data['desatualizado'] is True and data['valor']==165827
    with sf() as session:
        failed=session.scalar(select(Coleta).order_by(Coleta.id.desc()));assert failed.valida is False
        assert not session.scalar(select(func.count()).select_from(Observacao).where(Observacao.coleta_id==failed.id))


def test_o3_caged_no_collection_unavailable_and_simulation_isolated(sf,client):
    data=client.get('/api/v1/indicadores/caged-saldo').json();assert not data['disponivel'] and data['motivo'] and data['pontos']==[]
    ingest_caged(sf)
    with sf() as session:
        data=read_series(session,'caged-saldo','5a',test_mode=False).model_dump();assert not data['disponivel'] and data['pontos']==[]


def test_o3_caged_retry_after_and_limit(sf):
    calls=[];delays=[];metadata,raw=entry('2026-08')
    async def sleep(delay):delays.append(delay)
    def handler(request):
        calls.append(request);return httpx.Response(503,headers={'Retry-After':'999'}) if len(calls)<3 else httpx.Response(200,content=raw)
    async def run():
        module=importlib.import_module('backend.oil.caged')
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await module.CagedCollector(sf,client,simulada=True,sleep=sleep).collect(metadata)
    asyncio.run(run());assert len(calls)==3 and delays==[60,60]
    with pytest.raises(SourceFailure):ingest_caged(sf,handler=lambda r:httpx.Response(200,headers={'Content-Length':str(101*1024*1024)},content=b''))


def test_o3_caged_calendar_uses_end_of_day_brt(sf):
    module=importlib.import_module('backend.oil.caged');module.register_calendar(sf)
    with sf() as session:
        assert read_series(session,'caged-saldo','5a',test_mode=True,now=datetime(2026,10,30,3,0,tzinfo=timezone.utc)).atrasado_calendario is True
        assert read_series(session,'caged-saldo','5a',test_mode=True,now=datetime(2026,10,29,12,0,tzinfo=timezone.utc)).atrasado_calendario is False


def test_o3_migration_has_single_head_up_down_up(sf,engine):
    from alembic.script import ScriptDirectory
    assert ScriptDirectory.from_config(alembic_config()).get_heads()==['0004']
    command.downgrade(alembic_config(),'0003');command.upgrade(alembic_config(),'0004')
