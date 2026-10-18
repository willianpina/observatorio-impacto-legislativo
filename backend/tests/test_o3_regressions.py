"""Casos de isolamento, auditoria e chegada intermediária fora de ordem."""
import asyncio
import copy
import hashlib
import io

import httpx
import openpyxl
import pytest
from sqlalchemy import select

from backend.oil.caged import CagedCollector, parse_workbook, reconstruct
from backend.oil.ingest import SourceFailure
from backend.oil.legislation import LawCollector
from backend.oil.models import Coleta
from .conftest import no_sleep
from .test_o3_caged import entry, ingest_caged
from .test_o3_law import CAMARA, SENADO, ingest_law, payload, replay_law


def synthetic_file(sf,month,value):
    metadata,original=entry('2026-08')
    workbook=openpyxl.load_workbook(io.BytesIO(original));sheet=workbook['Tabela 5.1']
    sheet.cell(19,6,value)
    sheet.insert_rows(21);sheet.cell(21,2,'Setembro/2026');sheet.cell(21,6,1)
    stream=io.BytesIO();workbook.save(stream);raw=stream.getvalue()
    metadata.update(competencia=month,sha256=hashlib.sha256(raw).hexdigest())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r:httpx.Response(200,content=raw))) as client:
            return await CagedCollector(sf,client,simulada=True,sleep=no_sleep).collect(metadata)
    return asyncio.run(run())


def test_intermediate_file_cannot_change_future_unchanged_value(sf):
    ingest_caged(sf,'2026-07')
    synthetic_file(sf,'2026-09',58568)  # Julho igual ao arquivo inicial.
    with sf() as session:before=reconstruct(session,'2026-09',test_mode=True)
    ingest_caged(sf,'2026-08')  # Revisão intermediária recebida tarde.
    with sf() as session:
        after=reconstruct(session,'2026-09',test_mode=True)
        assert {p:v['valor'] for p,v in after.items()}=={p:v['valor'] for p,v in before.items()}
        assert reconstruct(session,'2026-08',test_mode=True)['202607']['valor']==55915


def test_invalid_collection_never_enters_reconstruction(sf):
    ingest_caged(sf)
    with pytest.raises(SourceFailure):ingest_caged(sf,raw=b'PLANILHA INVALIDA SIMULADA')
    with sf() as session:
        points=reconstruct(session,'2026-08',test_mode=True)
        failed=session.scalar(select(Coleta).order_by(Coleta.id.desc()))
        assert not failed.valida and all(p['coleta_id']!=failed.id for p in points.values())


def test_stream_limit_enforced_without_content_length(sf,monkeypatch):
    from backend.oil import caged
    assert caged.MAX_BYTES==100*1024*1024
    monkeypatch.setattr(caged,'MAX_BYTES',1024)
    class Stream(httpx.AsyncByteStream):
        async def __aiter__(self):
            for _ in range(3):yield b'x'*512
    with pytest.raises(SourceFailure):
        ingest_caged(sf,handler=lambda r:httpx.Response(200,stream=Stream()))
    with sf() as session:
        failed=session.scalar(select(Coleta).order_by(Coleta.id.desc()))
        assert not failed.valida and failed.resposta_bruta is None


@pytest.mark.parametrize('mutation',['title','gap','start','competencia'])
def test_strict_workbook_rejects_unexpected_source_layout(mutation):
    _,raw=entry('2026-08');wb=openpyxl.load_workbook(io.BytesIO(raw));ws=wb['Tabela 5.1']
    if mutation=='title':ws.cell(2,2,'TABELA 5 - SEM AJUSTES')
    if mutation=='gap':ws.cell(7,2,'Agosto/2025')
    stream=io.BytesIO();wb.save(stream)
    with pytest.raises(ValueError):
        parse_workbook(stream.getvalue(),'2026-07' if mutation=='competencia' else '2026-08',allow_recorte=mutation!='start')


def test_simulation_requires_explicit_test_mode(sf,monkeypatch):
    monkeypatch.delenv('OIL_TEST_MODE')
    for collector in (LawCollector,CagedCollector):
        with pytest.raises(ValueError):collector(sf,None,simulada=True)


@pytest.mark.parametrize('house',['camara','senado'])
def test_invalid_json_is_audited_without_snapshot_for_house(sf,client,house):
    def handler(request):
        if str(request.url)==(CAMARA if house=='camara' else SENADO):return httpx.Response(200,text='JSON INVALIDO SIMULADO')
        return replay_law()(request)
    ingest_law(sf,handler)
    data=client.get('/api/v1/proposicoes/2233802').json()
    assert next(s for s in data['situacao'] if s['casa']==house)['disponivel'] is False
    with sf() as session:
        failed=session.scalar(select(Coleta).where(Coleta.serie_codigo==f'proposicao:2233802:{house}'))
        assert failed.erro=='resposta_invalida' and not failed.valida


def test_law_cached_house_is_preserved_and_flagged(sf,client):
    ingest_law(sf);ingest_law(sf,replay_law(fail='senado'))
    data=client.get('/api/v1/proposicoes/2233802').json()
    assert data['desatualizado'] is True
    senate=next(s for s in data['situacao'] if s['casa']=='senado')
    assert senate['coleta_id'] and not senate['disponivel']
    assert senate['situacao']==payload('senado')['situacaoAtual']


def test_law_invalid_generated_norm_is_not_validated(sf,client):
    altered=copy.deepcopy(payload('senado'));altered['normaGerada']={'rotulo':'INCOMPLETO SIMULADO'}
    ingest_law(sf,replay_law(senado=altered))
    data=client.get('/api/v1/proposicoes/2233802').json()
    assert next(s for s in data['situacao'] if s['casa']=='senado')['disponivel'] is False


def test_law_retry_after_respected(sf):
    calls=0;delays=[]
    async def sleep(value):delays.append(value)
    def handler(request):
        nonlocal calls
        if str(request.url)==SENADO:
            calls+=1
            if calls<3:return httpx.Response(429,headers={'Retry-After':'999'})
        return replay_law()(request)
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            await LawCollector(sf,client,simulada=True,sleep=sleep).collect(2233802)
    asyncio.run(run());assert delays==[60,60]


def test_law_identification_comes_from_selected_snapshot_not_shared_catalog(sf,client):
    from backend.oil.models import Proposicao
    ingest_law(sf)
    with sf.begin() as session:
        prop=session.scalar(select(Proposicao));prop.sigla_tipo='TESTE';prop.numero=1
    data=client.get('/api/v1/proposicoes/2233802').json()
    assert data['sigla_tipo']==payload('camara')['dados']['siglaTipo']
    assert data['numero']==payload('camara')['dados']['numero']
