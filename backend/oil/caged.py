"""Leitura estrita da Tabela 5.1 e reconstrução por competência do arquivo."""
import asyncio
import hashlib
import io
import logging
import math
import os
import re
import tempfile
from urllib.parse import urlsplit, parse_qs
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

import httpx
import openpyxl
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert

from .ingest import SourceFailure, ensure_unavailable_catalog, retry_delay
from .models import CagedArquivo, Coleta, Divulgacao, Fonte, Observacao, Serie
from .schemas import SeriesResponse

MTE='https://www.gov.br/trabalho-e-emprego/pt-br/acesso-a-informacao/acoes-e-programas/programas-projetos-acoes-obras-e-atividades/estatisticas-trabalho'
MONTHS=('Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro')
MAX_BYTES=100*1024*1024
MAX_EXPANDED_BYTES=512*1024*1024
MAX_COMPRESSION_RATIO=100
MAX_DATA_ROWS=1000
LOG=logging.getLogger(__name__)


def download_timeout():
    """Configuração exclusiva do job; nunca desativa o prazo nem muda o cliente comum."""
    try:
        read=float(os.environ.get('OIL_CAGED_READ_TIMEOUT','120'))
    except ValueError:
        raise ValueError('OIL_CAGED_READ_TIMEOUT deve ser um número positivo e finito.') from None
    if not math.isfinite(read) or read<=0:
        raise ValueError('OIL_CAGED_READ_TIMEOUT deve ser um número positivo e finito.')
    return httpx.Timeout(connect=5,read=min(read,300),write=10,pool=10)


_SUBCODES={
    'Conteúdo descompactado acima do limite.':'zip_total',
    'Razão de compactação acima do limite.':'zip_razao',
    'Aba esperada ausente.':'aba',
    'Cabeçalho inesperado.':'cabecalho',
    'Título inesperado.':'titulo',
    'Rodapé inesperado.':'rodape',
    'Série deve iniciar em janeiro de 2020.':'inicio',
    'Tabela acima do limite de linhas.':'linhas',
    'Mês inválido.':'mes',
    'Saldo não inteiro.':'saldo',
    'Meses descontínuos.':'continuidade',
    'Último mês divergente.':'competencia',
    'Competência inválida.':'competencia',
    'Hash divergente do manifesto.':'hash',
    'Arquivo acima do limite.':'download',
    'Arquivo acima do limite ou tamanho inválido.':'download',
}


def failure_subcode(exc):
    if isinstance(exc,BadZipFile):return 'zip'
    if isinstance(exc,(ParseError,DefusedXmlException)):return 'xml'
    if isinstance(exc,openpyxl.utils.exceptions.InvalidFileException):return 'arquivo'
    if isinstance(exc,KeyError):return 'chave'
    if isinstance(exc,TypeError):return 'tipo'
    return _SUBCODES.get(str(exc),'parser')


def parse_workbook(source, competencia, *, allow_recorte=False):
    if not re.fullmatch(r'[0-9]{4}-(0[1-9]|1[0-2])',competencia):
        raise ValueError('Competência inválida.')
    stream=io.BytesIO(source) if isinstance(source,bytes) else source
    with ZipFile(stream) as archive:
        members=archive.infolist()
        if sum(member.file_size for member in members)>MAX_EXPANDED_BYTES:
            raise ValueError('Conteúdo descompactado acima do limite.')
        for member in members:
            if member.file_size and (member.compress_size==0 or member.file_size/member.compress_size>MAX_COMPRESSION_RATIO):
                raise ValueError('Razão de compactação acima do limite.')
    if hasattr(stream,'seek'):stream.seek(0)
    workbook=openpyxl.load_workbook(stream,read_only=True,data_only=True)
    try:
        if 'Tabela 5.1' not in workbook.sheetnames: raise ValueError('Aba esperada ausente.')
        sheet=workbook['Tabela 5.1']
        title=sheet.cell(2,2).value
        if not isinstance(title,str) or not title.startswith('TABELA 5.1') or 'COM AJUSTES' not in title:
            raise ValueError('Título inesperado.')
        header=tuple(next(sheet.iter_rows(min_row=5,max_row=5,min_col=2,max_col=7,values_only=True)))
        if header!=('Mês','Estoque','Admissões','Desligamentos','Saldos','Variação Relativa (%)'):
            raise ValueError('Cabeçalho inesperado.')
        points={};ended=False
        for index,row in enumerate(sheet.iter_rows(min_row=6,max_row=6+MAX_DATA_ROWS,min_col=2,max_col=6,values_only=True)):
            name,value=row[0],row[4]
            if name is None:
                ended=True;continue
            if ended:
                if not isinstance(name,str) or not name.startswith(('Fonte','*')): raise ValueError('Rodapé inesperado.')
                continue
            if index>=MAX_DATA_ROWS:
                raise ValueError('Tabela acima do limite de linhas.')
            if not isinstance(name,str) or not re.fullmatch(r'[^/]+/[0-9]{4}',name): raise ValueError('Mês inválido.')
            month,year=name.split('/')
            if month not in MONTHS: raise ValueError('Mês inválido.')
            period=year+f'{MONTHS.index(month)+1:02d}'
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not float(value).is_integer():
                raise ValueError('Saldo não inteiro.')
            if points:
                prev=max(points); expected=int(prev[:4])*12+int(prev[4:])
                if int(year)*12+MONTHS.index(month)!=expected: raise ValueError('Meses descontínuos.')
            elif period!='202001' and not allow_recorte:
                raise ValueError('Série deve iniciar em janeiro de 2020.')
            points[period]=int(value)
        if not points or max(points)!=competencia.replace('-',''): raise ValueError('Último mês divergente.')
        return points
    finally:
        workbook.close()


def good_collections(session, test_mode):
    return select(Coleta).where(Coleta.serie_codigo=='caged-saldo',Coleta.simulada==test_mode,Coleta.valida.is_(True))


def reconstruct(session, competencia, *, test_mode=False):
    target=session.scalar(good_collections(session,test_mode).where(Coleta.mes_referencia==competencia)
        .order_by(Coleta.id.desc()).limit(1))
    if target is None: return {}
    sid=session.scalar(select(Serie.id).where(Serie.codigo=='caged-saldo'))
    records=session.execute(select(Observacao,Coleta).join(Coleta).where(Observacao.serie_id==sid,
        Coleta.simulada==test_mode,Coleta.valida.is_(True),Coleta.mes_referencia<=competencia,
        Observacao.periodo>=target.periodo_inicio,Observacao.periodo<=target.periodo_fim)
        .order_by(Coleta.mes_referencia,Observacao.versao)).all()
    return {o.periodo:dict(valor=int(o.valor) if o.valor is not None else None,coleta_id=c.id,
        versao=o.versao,coletado_em=c.coletado_em) for o,c in records}


class CagedCollector:
    def __init__(self,sessions,client,*,simulada=False,sleep=asyncio.sleep):
        if simulada and os.environ.get('OIL_TEST_MODE')!='1':
            raise ValueError('Coleta simulada exige modo de teste explícito.')
        self.sessions,self.client,self.simulada,self.sleep=sessions,client,simulada,sleep

    def record(self,entry,*,status=None,error=None,digest=None,size=None):
        with self.sessions.begin() as session:
            ensure_unavailable_catalog(session)
            source=session.scalar(select(Fonte).where(Fonte.codigo=='mte-caged'))
            source.link=MTE+'/novo-caged'
            item=Coleta(fonte_id=source.id,serie_codigo='caged-saldo',url_consultada=entry['url_download'],
                coletado_em=datetime.now(timezone.utc),codigo_http=status,hash_sha256=digest,
                resposta_bruta=None,simulada=self.simulada,erro=error,tipo='arquivo',valida=False,
                mes_referencia=entry['competencia'],tamanho_bytes=size,
                manifesto_json={key:entry.get(key) for key in ('arquivo','drive_file_id','url_download','pagina_oficial','publicado_em')})
            session.add(item);session.flush();return item.id

    async def collect(self,entry):
        # A exceção para recortes nunca pode ser ativada em uma coleta de produção.
        recorte=bool(entry.get('fixture_recorte'))
        if recorte and not (self.simulada and os.environ.get('OIL_TEST_MODE')=='1'):
            raise ValueError('Recorte permitido somente em teste explícito.')
        expected_hash=entry.get('sha256')
        if not isinstance(expected_hash,str) or re.fullmatch(r'[0-9a-fA-F]{64}',expected_hash) is None:
            self.record(entry,error='arquivo_invalido')
            raise SourceFailure('arquivo_invalido')
        if not re.fullmatch(r'[0-9]{4}-(0[1-9]|1[0-2])',entry['competencia']):
            raise ValueError('Competência inválida no manifesto.')
        page=urlsplit(entry['pagina_oficial'])
        if page.scheme!='https' or page.netloc!='www.gov.br' or not page.path.startswith('/trabalho-e-emprego/'):
            raise ValueError('Página de atribuição fora da fonte autorizada.')
        url=entry['url_download']
        if not re.fullmatch(r'https://drive\.usercontent\.google\.com/download\?[^\s]+',url):
            raise ValueError('URL fora da origem autorizada no manifesto.')
        if not self.simulada and parse_qs(urlsplit(url).query).get('id')!=[entry['drive_file_id']]:
            raise ValueError('Identificador de arquivo divergente do manifesto.')
        url=str(httpx.URL(url).copy_set_param('confirm','t'))
        timeout=download_timeout()
        for attempt in range(3):
            response=None;cid=None;digest=hashlib.sha256();size=0
            with tempfile.TemporaryFile() as stream:
                try:
                    async with self.client.stream('GET',url,timeout=timeout,follow_redirects=False,
                        headers={'Accept':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}) as response:
                        if response.status_code!=200:
                            failure=SourceFailure('http_fonte',503 if response.status_code in {429,500,502,503,504} else 502)
                            cid=self.record(entry,status=response.status_code,error=failure.code)
                            if failure.status!=503: raise failure
                        else:
                            length=response.headers.get('Content-Length')
                            if length and (not length.isdecimal() or int(length)>MAX_BYTES):
                                raise ValueError('Arquivo acima do limite ou tamanho inválido.')
                            async for chunk in response.aiter_bytes():
                                size+=len(chunk)
                                digest.update(chunk)
                                if size>MAX_BYTES: raise ValueError('Arquivo acima do limite.')
                                stream.write(chunk)
                            cid=self.record(entry,status=200,digest=digest.hexdigest(),size=size)
                            if digest.hexdigest()!=expected_hash.lower():
                                raise ValueError('Hash divergente do manifesto.')
                            stream.seek(0)
                            points=parse_workbook(stream,entry['competencia'],allow_recorte=recorte)
                            count=self.persist(entry,points,cid,digest.hexdigest())
                            return dict(coleta_id=cid,novas_versoes=count,simulada=self.simulada)
                except httpx.TimeoutException as exc:
                    failure=SourceFailure('timeout_fonte',504)
                    # Escrita é o envio da requisição (write=10), não a conexão nem a espera do primeiro byte.
                    if isinstance(exc,(httpx.ConnectTimeout,httpx.PoolTimeout)): phase='timeout_conexao'
                    elif isinstance(exc,httpx.WriteTimeout): phase='timeout_escrita'
                    elif size: phase='timeout_leitura'
                    else: phase='timeout_primeiro_byte'
                    self.record(entry,status=response.status_code if response is not None else None,
                        error=phase,digest=digest.hexdigest() if size else None,size=size)
                except httpx.TransportError:
                    failure=SourceFailure('erro_conexao',503)
                    self.record(entry,error=failure.code)
                except (ValueError,KeyError,TypeError,BadZipFile,ParseError,DefusedXmlException,openpyxl.utils.exceptions.InvalidFileException) as exc:
                    subcode=failure_subcode(exc)
                    LOG.warning('caged arquivo_invalido subcodigo=%s',subcode)
                    if cid is None: cid=self.record(entry,status=response.status_code if response is not None else None,
                        error=subcode,digest=digest.hexdigest() if size else None,size=size)
                    else:
                        with self.sessions.begin() as session: session.get(Coleta,cid).erro=subcode
                    raise SourceFailure('arquivo_invalido') from None
            if attempt<2: await self.sleep(retry_delay(response,attempt))
        raise failure

    def persist(self,entry,points,cid,digest):
        with self.sessions.begin() as session:
            serie=session.scalar(select(Serie).where(Serie.codigo=='caged-saldo').with_for_update())
            collection=session.get(Coleta,cid)
            collection.periodo_inicio=min(points);collection.periodo_fim=max(points)
            future=session.scalars(good_collections(session,self.simulada).where(
                Coleta.mes_referencia>entry['competencia']).order_by(Coleta.mes_referencia,Coleta.id)).all()
            preserved={c.mes_referencia:reconstruct(session,c.mes_referencia,test_mode=self.simulada) for c in future}
            # Competência da fonte define precedência; versão é uma trilha de ingestão.
            previous=session.execute(select(Observacao,Coleta).join(Coleta).where(
                Observacao.serie_id==serie.id,Coleta.simulada==self.simulada,Coleta.valida.is_(True),
                Coleta.mes_referencia<=entry['competencia']).order_by(Coleta.mes_referencia,Observacao.versao)).all()
            latest={o.periodo:o for o,c in previous}
            versions={}
            for o in session.scalars(select(Observacao).where(Observacao.serie_id==serie.id)):
                versions[o.periodo]=max(versions.get(o.periodo,0),o.versao)
            count=0
            for period,value in points.items():
                if period in latest and latest[period].valor==value: continue
                session.add(Observacao(serie_id=serie.id,periodo=period,valor=value,cv=None,
                    variacao_anual=None,diferenca_anual=None,classificacao=None,classificacao_origem=None,
                    variacao_anual_origem='calculo_observatorio',coleta_id=cid,versao=versions.get(period,0)+1))
                count+=1
                versions[period]=versions.get(period,0)+1
            collection.valida=True;serie.disponivel=True;serie.motivo=None
            session.flush()
            # Um arquivo intermediário recebido tarde não pode mudar um arquivo
            # futuro já validado. Materializa a fronteira que antes era implícita.
            for month,expected in preserved.items():
                current=reconstruct(session,month,test_mode=self.simulada)
                boundary=next(c for c in reversed(future) if c.mes_referencia==month)
                for period,item in expected.items():
                    if current.get(period,{}).get('valor')==item['valor']:continue
                    versions[period]=versions.get(period,0)+1
                    session.add(Observacao(serie_id=serie.id,periodo=period,valor=item['valor'],cv=None,
                        variacao_anual=None,diferenca_anual=None,classificacao=None,classificacao_origem=None,
                        variacao_anual_origem='calculo_observatorio',coleta_id=boundary.id,versao=versions[period]))
                    count+=1
                session.flush()
            serie.tipo_valor=None;serie.unidade=serie.unidade_variacao='postos de trabalho'
            values={k:entry.get(k) for k in ('competencia','arquivo','drive_file_id','url_download','pagina_oficial','publicado_em')}
            values.update(sha256=digest,coleta_id=cid)
            statement=insert(CagedArquivo).values(**values)
            session.execute(statement.on_conflict_do_update(index_elements=['competencia'],set_=values))
            return count


def register_calendar(sessions):
    with sessions.begin() as session:
        ensure_unavailable_catalog(session)
        sid=session.scalar(select(Serie.id).where(Serie.codigo=='caged-saldo'))
        for period,date in [('202609','2026-10-29'),('202610','2026-11-30'),('202611','2026-12-30')]:
            values=dict(serie_id=sid,periodo=period,prevista_em=datetime.fromisoformat(date+'T23:59:00-03:00'),
                fonte='MTE — calendário público do Novo Caged',link=MTE+'/o-pdet/calendario-de-divulgacao-do-novo-caged')
            session.execute(insert(Divulgacao).values(**values).on_conflict_do_nothing(constraint='uq_divulgacao_periodo'))


def read_caged(session,janela,*,test_mode=False,now=None):
    from .read import read_unavailable,window_start,periods_between
    now=now or datetime.now(timezone.utc)
    good=session.scalar(good_collections(session,test_mode).order_by(Coleta.mes_referencia.desc(),Coleta.id.desc()).limit(1))
    serie=session.scalar(select(Serie).where(Serie.codigo=='caged-saldo'))
    points=reconstruct(session,good.mes_referencia,test_mode=test_mode) if good else {}
    due=session.scalars(select(Divulgacao).where(Divulgacao.serie_id==serie.id,Divulgacao.prevista_em<now)).all() if serie else []
    delayed=any(item.periodo not in points for item in due)
    if not good:
        result=read_unavailable(session,janela)
        result.atrasado_calendario=delayed
        result.motivo='Nenhuma coleta válida do Novo Caged disponível.'
        return result
    metadata=good.manifesto_json
    failed=session.scalar(select(Coleta).where(Coleta.serie_codigo=='caged-saldo',Coleta.simulada==test_mode,
        Coleta.erro.is_not(None)).order_by(Coleta.id.desc()).limit(1))
    month=MONTHS[int(good.mes_referencia[5:])-1].lower()+'/'+good.mes_referencia[:4]
    common=dict(tipo_valor=None,tipo_periodo='mensal',medida='saldo',unidade='postos de trabalho',
        fonte='MTE / Novo Caged',tabela_ou_arquivo=metadata['arquivo'],link=metadata['pagina_oficial'],
        cv=None,classificacao=None,classificacao_origem=None,variacao_anual_unidade='postos de trabalho',
        variacao_anual_origem='calculo_observatorio',ajuste='com_ajuste',competencia_coleta=good.mes_referencia,
        atribuicao=dict(texto=f'Fonte: Novo Caged – MTE, Tabela 5.1 (com ajustes), arquivo {month}, coletado em {good.coletado_em.astimezone(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y")}. Elaboração: OIL.',
            link=metadata['pagina_oficial'],licenca=None,base_legal=dict(
                texto='Uso livre e gratuito, com citação da fonte (Decreto nº 8.777/2016, art. 4º e §1º, com a redação do Decreto nº 9.903/2019). O MTE não publica licença própria para estes dados.',
                link='https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2016/decreto/d8777.htm')),
        nota_revisao='Série revisada mensalmente pelo MTE: os meses anteriores podem mudar a cada novo arquivo.',
        desatualizado=failed is not None and failed.id>good.id,atrasado_calendario=delayed,simulada=test_mode,
        variaveis={},disponivel=True,motivo=None)
    rows=[];start=max(min(points),window_start(max(points),janela))
    for period in periods_between(start,max(points)):
        item=points.get(period);value=item['valor'] if item else None
        previous=points.get(str(int(period[:4])-1)+period[4:])
        difference=value-previous['valor'] if value is not None and previous and previous['valor'] is not None else None
        rows.append(dict(common,periodo=period,valor=value,lacuna=value is None,variacao_anual=difference,
            coleta_id=item['coleta_id'] if item else None,coletado_em=item['coletado_em'] if item else None,
            versao=item['versao'] if item else None))
    last=rows[-1]
    return SeriesResponse(**common,codigo='caged-saldo',nome=serie.nome,janela=janela,pontos=rows,
        valor=last['valor'],periodo=last['periodo'],lacuna=last['lacuna'],variacao_anual=last['variacao_anual'],
        coleta_id=good.id,coletado_em=good.coletado_em)
